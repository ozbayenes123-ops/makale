# -*- coding: utf-8 -*-
"""Makale formatında Word derleme: yapılandırılmış _tr.txt -> biçimli .docx.

Özellikler (hepsi `config.json > docx` ile ayarlanabilir):
- A4 sayfa, yapılandırılabilir kenar boşlukları
- Gövde: Times New Roman 12, iki yana yaslı, 1.5 satır aralığı, ilk satır girintisi
- Başlık/alt başlık bloğu, H1/H2 stilleri (siyah, akademik)
- GERÇEK Word dipnotları: metindeki [fn N] göndermeleri Word dipnotuna
  dönüşür (sayfa altında numaralı, üstsimge göndermeli)
- İçindekiler (TOC alanı — Word açılışta günceller), sayfa numaraları
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as open_docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

from makale_pipeline.config import DEFAULT_DOCX
from makale_pipeline.paths import find_source_for_target
from makale_pipeline.structured import (
    FN_REF_RE,
    parse_file,
    source_contains_toc_heading,
)

_FOOTNOTE_MODES = ("auto", "on", "off")

_FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument"
    ".wordprocessingml.footnotes+xml"
)
_FOOTNOTES_RT = (
    "http://schemas.openxmlformats.org/officeDocument/2006"
    "/relationships/footnotes"
)
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_ALIGN = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _merge_style(override: dict | None) -> dict:
    merged = dict(DEFAULT_DOCX)
    if override:
        for key, value in override.items():
            if isinstance(value, dict) and isinstance(merged.get(key), dict):
                merged[key] = {**merged[key], **value}
            else:
                merged[key] = value
    return merged


def resolve_footnote_mode(style: dict) -> str:
    """Dipnot kipini çözer: 'auto' | 'on' | 'off' (bool de kabul edilir)."""
    mode = style.get("footnotes", "auto")
    if isinstance(mode, bool):
        return "on" if mode else "off"
    mode = str(mode).strip().lower()
    return mode if mode in _FOOTNOTE_MODES else "auto"


def resolve_toc_decision(
    style: dict, doc_model, translation_path: Path
) -> tuple[bool, str]:
    """İçindekiler kararını verir.

    docx.toc=true/false zorlar; "auto" ise yalnızca yapılandırılmış taslakta
    [TOC] girişleri veya kaynak dosyada gerçek bir içindekiler başlığı varsa
    eklenir. Dönen: (eklensin_mi, kip).
    """
    raw = style.get("toc", "auto")
    if isinstance(raw, bool):
        return raw, "on" if raw else "off"
    if isinstance(raw, str):
        low = raw.strip().lower()
        if low in ("true", "on", "yes", "1"):
            return True, "on"
        if low in ("false", "off", "no", "0"):
            return False, "off"
    # "auto" (veya tanınmayan değer): varlığa bak.
    if doc_model.toc:
        return True, "auto"
    source = find_source_for_target(Path(translation_path))
    if source is not None:
        try:
            if source_contains_toc_heading(source.read_text(encoding="utf-8")):
                return True, "auto"
        except OSError:
            pass
    return False, "auto"


def _strip_fn_refs(text: str) -> str:
    """Metinden [fn N] göndermelerini temiz biçimde siler (dipnotsuz sürüm)."""
    cleaned = FN_REF_RE.sub("", text)
    cleaned = re.sub(r"\s+([,.;:!?])", r"\1", cleaned)  # boşluk+noktalama
    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned).strip()
    return cleaned


def _force_font(style, font: str) -> None:
    """Stil yazı tipini tema bağımsız sabitler.

    `style.font.name = X` yalnızca w:ascii/w:hAnsi yazar; şablondan gelen
    w:asciiTheme="minorHAnsi" özniteliği silinmezse Word tema yazı tipini
    (Calibri) kullanır ve başlık/alt başlık TNR görünmez. Ölçüldü: Word'de
    Title/Heading paragrafları Calibri çıkıyordu.
    """
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)


def _force_doc_default_font(doc, font: str) -> None:
    """docDefaults rPrDefault yazı tipini de sabitler (kalıtılan her şey için)."""
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)


def _normalize_all_fonts(doc, font: str) -> None:
    """Şablondaki dokunulmamış her stilde tema özniteliğini siler.

    Yalnızca düzenlediğimiz stillere bakmak yetmez: python-docx şablonundaki
    diğer stiller (Hyperlink, List Paragraph, Quote, ...) hâlâ
    `asciiTheme="minorHAnsi"` taşır ve Word o alanlarda tema yazı tipini
    (Calibri/Cambria) kullanır. Test, styles.xml genelinde tema özniteliği
    kalmamasını şart koşar.
    """
    styles_el = doc.styles.element
    for rfonts in styles_el.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            val = rfonts.get(qn(attr))
            if val is None or val in ("Cambria", "Calibri", "Calibri Light"):
                rfonts.set(qn(attr), font)


def _ensure_styles(doc, style: dict):
    """Normal/Title/Heading/Subtitle/Footnote stillerini makale biçimine sokar."""
    body_font = style["body_font"]
    body_size = Pt(style["body_size_pt"])

    _force_doc_default_font(doc, body_font)

    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = body_size
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = _ALIGN.get(style["body_alignment"], WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf.line_spacing = style["line_spacing"]
    pf.space_after = Pt(style["space_after_pt"])
    pf.first_line_indent = Cm(style["first_line_indent_cm"])
    _force_font(normal, body_font)

    title = doc.styles["Title"]
    title.font.name = body_font
    title.font.size = Pt(style["title_size_pt"])
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    _force_font(title, body_font)

    try:
        subtitle = doc.styles["Subtitle"]
    except KeyError:
        subtitle = None
    if subtitle is not None:
        subtitle.font.name = body_font
        subtitle.font.size = Pt(style["subtitle_size_pt"])
        subtitle.font.bold = False
        subtitle.font.italic = True
        subtitle.font.color.rgb = RGBColor(0, 0, 0)
        _force_font(subtitle, body_font)

    for name, size_key, bold, italic in (
        ("Heading 1", "heading1_size_pt", True, False),
        ("Heading 2", "heading2_size_pt", True, False),
    ):
        hs = doc.styles[name]
        hs.font.name = body_font
        hs.font.size = Pt(style[size_key])
        hs.font.bold = bold
        hs.font.italic = italic
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 6)
        hs.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 3)
        hs.paragraph_format.keep_with_next = True
        _force_font(hs, body_font)

    try:
        fn_text = doc.styles["Footnote Text"]
    except KeyError:
        fn_text = doc.styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
    fn_text.font.name = body_font
    fn_text.font.size = Pt(style["footnote_size_pt"])
    fn_text.font.color.rgb = RGBColor(0, 0, 0)
    fn_text.paragraph_format.space_after = Pt(2)
    fn_text.paragraph_format.line_spacing = 1.0
    _force_font(fn_text, body_font)

    try:
        fn_ref = doc.styles["Footnote Reference"]
    except KeyError:
        fn_ref = doc.styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
    fn_ref.font.superscript = True
    _force_font(fn_ref, body_font)

    _normalize_all_fonts(doc, body_font)


def _add_field(paragraph, instr: str):
    """Paragrafa alan kodu ekler (TOC, PAGE vb.)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run2 = paragraph.add_run()
    run2._r.append(instr_el)
    run3 = paragraph.add_run()
    run3._r.append(fld_separate)
    run4 = paragraph.add_run("  (Word'de alanı güncelleyin: F9)  ")
    run4.italic = True
    run5 = paragraph.add_run()
    run5._r.append(fld_end)


def _set_update_fields(doc):
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


def _footnote_id(num: str, fallback: int) -> int:
    try:
        fid = int(num)
        return fid if fid > 0 else fallback
    except (TypeError, ValueError):
        return fallback


def _build_footnotes_part(doc, footnotes: list, style: dict) -> dict[str, int]:
    """word/footnotes.xml bölümünü kurar; dipnot numarası -> w:id eşlemesi."""
    size_half = int(style["footnote_size_pt"] * 2)
    font = style["body_font"]
    root = OxmlElement("w:footnotes")

    for special, stype in (("-1", "separator"), ("0", "continuationSeparator")):
        fn = OxmlElement("w:footnote")
        fn.set(qn("w:id"), special)
        p = OxmlElement("w:p")
        r = OxmlElement("w:r")
        sc = OxmlElement("w:separator") if stype == "separator" else OxmlElement("w:continuationSeparator")
        r.append(sc)
        p.append(r)
        fn.append(p)
        root.append(fn)

    mapping: dict[str, int] = {}
    next_id = 1
    for fn in footnotes:
        fid = _footnote_id(fn.num, 1000 + next_id)
        while fid in mapping.values():
            next_id += 1
            fid = 1000 + next_id
        mapping[str(fn.num)] = fid
        next_id += 1
        fel = OxmlElement("w:footnote")
        fel.set(qn("w:id"), str(fid))
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), "FootnoteText")
        ppr.append(pstyle)
        p.append(ppr)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), font)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size_half))
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), str(size_half))
        rpr.append(rfonts)
        rpr.append(sz)
        rpr.append(sz_cs)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = fn.text
        r.append(t)
        p.append(r)
        fel.append(p)
        root.append(fel)

    xml_bytes = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    from docx.opc.part import Part

    partname = PackURI("/word/footnotes.xml")
    part = Part(partname, _FOOTNOTES_CT, xml_bytes, doc.part.package)
    doc.part.relate_to(part, _FOOTNOTES_RT)
    return mapping


def _add_paragraph_with_footnotes(doc, text: str, mapping: dict[str, int]):
    """[fn N] göndermelerini gerçek Word dipnotlarına çevirerek paragraf ekler."""
    para = doc.add_paragraph()
    parts = FN_REF_RE.split(text)
    for i, chunk in enumerate(parts):
        if i % 2 == 0:
            if chunk:
                para.add_run(chunk)
        else:
            fid = mapping.get(chunk)
            if fid is None:
                para.add_run(f"[{chunk}]")
                continue
            run = para.add_run()
            rstyle = run._r.get_or_add_rPr()
            rs = OxmlElement("w:rStyle")
            rs.set(qn("w:val"), "FootnoteReference")
            rstyle.append(rs)
            fr = OxmlElement("w:footnoteReference")
            fr.set(qn("w:id"), str(fid))
            run._r.append(fr)
    return para


def export_article_docx(
    translation_path: Path | str,
    output_path: Path | str | None = None,
    style: dict | None = None,
) -> dict:
    """Yapılandırılmış _tr.txt dosyasını makale formatında .docx'e derler."""
    tr = Path(translation_path)
    if not tr.exists():
        raise FileNotFoundError(f"Çeviri dosyası yok: {tr}")
    doc_model = parse_file(tr)
    if not doc_model.is_structured:
        raise ValueError(f"Yapılandırılmamış belge ([BODY] yok): {tr.name}")
    style = _merge_style(style)

    doc = open_docx()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    margins = style["margin_cm"]
    section.top_margin = Cm(margins["top"])
    section.bottom_margin = Cm(margins["bottom"])
    section.left_margin = Cm(margins["left"])
    section.right_margin = Cm(margins["right"])

    _ensure_styles(doc, style)

    footnote_mode = resolve_footnote_mode(style)
    warnings: list[str] = []
    if footnote_mode == "on" and doc_model.footnote_count == 0:
        warnings.append(
            "Dipnot kipi 'on' ama taslakta hiç dipnot yok ([FOOTNOTES]/[fn N] "
            "bulunamadı). Dipnotsuz derlendi."
        )

    keep_footnotes = footnote_mode in ("auto", "on")
    mapping: dict[str, int] = {}
    if keep_footnotes:
        mapping = _build_footnotes_part(doc, doc_model.footnotes, style)

    if doc_model.title:
        doc.add_paragraph(doc_model.title, style="Title")
    if doc_model.subtitle:
        p = doc.add_paragraph(style="Subtitle")
        run = p.add_run(doc_model.subtitle)
        run.font.size = Pt(style["subtitle_size_pt"])
        run.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if doc_model.vat_label:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(doc_model.vat_label)
        run.font.size = Pt(12)

    # Yazar / kurum: taslak bloğu, yoksa config yedeği (uydurma değer üretilmez).
    author = (doc_model.author or style.get("author") or "").strip()
    institution = (doc_model.institution or style.get("institution") or "").strip()
    if author:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(author)
    if institution:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(institution)
        run.italic = True

    include_toc, toc_mode = resolve_toc_decision(style, doc_model, tr)
    if include_toc:
        doc.add_heading(style.get("toc_title", "İçindekiler"), level=1)
        _add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')

    unreferenced: list[str] = []
    body_has_ref = doc_model.body_text()
    for sec in doc_model.sections:
        if sec.title and sec.title != "GÖVDE":
            doc.add_heading(sec.title, level=1)
        for sub in sec.subsections:
            doc.add_heading(sub.title, level=2)
        for para in sec.paragraphs:
            if not para.text:
                continue
            if keep_footnotes:
                _add_paragraph_with_footnotes(doc, para.text, mapping)
            else:
                text = _strip_fn_refs(para.text)
                if text:
                    doc.add_paragraph(text)

    if keep_footnotes:
        referenced = set(mapping.keys())
        for fn in doc_model.footnotes:
            if str(fn.num) not in referenced or f"[fn {fn.num}]" not in body_has_ref:
                unreferenced.append(fn.num)
        if unreferenced:
            doc.add_heading("Dipnotlar", level=1)
            table = {fn.num: fn.text for fn in doc_model.footnotes}
            for num in unreferenced:
                p = doc.add_paragraph()
                r = p.add_run(f"{num}. ")
                r.bold = True
                p.add_run(table.get(num, ""))

    if style.get("page_numbers"):
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(p, "PAGE")

    _set_update_fields(doc)

    out = Path(output_path) if output_path else tr.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return {
        "source": str(tr),
        "output": str(out),
        "paragraphs": doc_model.paragraph_count,
        "footnotes": doc_model.footnote_count,
        "real_footnotes": len(mapping),
        "footnote_mode": footnote_mode,
        "toc_mode": toc_mode,
        "toc_included": include_toc,
        "author": author,
        "institution": institution,
        "warnings": warnings,
        "format": "makale",
    }
